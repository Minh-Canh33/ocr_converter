import easyocr
import cv2
import numpy as np

import tkinter as tk
from tkinter import filedialog
from docx import Document


# Ẩn cửa sổ chính của tkinter
root = tk.Tk()
root.withdraw()

# Mở hộp thoại chọn file và lấy đường dẫn
file_path = filedialog.askopenfilename(
    title="Chọn file",
    filetypes=[("Image files", "*.jpg"),("Image files", "*.png"), ("All files", "*.*")]
)
if file_path:
    path = cv2.imdecode(np.fromfile(file_path,dtype=np.uint8),cv2.IMREAD_COLOR)
    # print(path)
    reader = easyocr.Reader(['vi','en'])
    result = reader.readtext(path)
    # with open("tam.txt",'w',encoding='utf-8') as f:
    #     for(bbox,text,prob) in result:
    #         f.write(text)
    s = ""
    for(bbox,text,prob) in result:
        print(text)
        print(bbox)
        s += "".join(text + " ")
    print("_"*20)
    print(s)

    # add tkinter GUI to display OCR result stored in variable s
    # add button to save as docx and pdf
    # reuse existing OCR logic, do not rewrite it
    # Create a new window to display results
    result_window = tk.Toplevel(root)
    result_window.title("OCR Result")
    result_window.geometry("600x500")

    # Text area to display the OCR result
    text_area = tk.Text(result_window, wrap=tk.WORD, font=("Arial", 12))
    text_area.insert(tk.END, s)
    text_area.pack(expand=True, fill='both', padx=10, pady=10)

    def save_as_docx():
        save_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Document", "*.docx")])
        if save_path:
            doc = Document()
            doc.add_paragraph(text_area.get("1.0", tk.END))
            doc.save(save_path)
            tk.messagebox.showinfo("Success", "Saved as DOCX successfully!")
    
    # def save_as_pdf():
    #     save_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Document", "*.pdf")])
    #     if save_path:
    #         from fpdf import FPDF
    #         pdf = FPDF()
    #         pdf.add_page()
    #         pdf.set_auto_page_break(auto=True, margin=15)
    #         pdf.set_font("Arial", size=12)
    #         for line in text_area.get("1.0", tk.END).splitlines():
    #             pdf.cell(0, 10, txt=line, ln=True)
    #         pdf.output(save_path)
    #         tk.messagebox.showinfo("Success", "Saved as PDF successfully!")

